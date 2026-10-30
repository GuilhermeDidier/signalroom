"""Small real PDF / DOCX files built on the fly for ingestion tests."""

from __future__ import annotations

import io

from docx import Document
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate


def make_pdf() -> bytes:
    styles = getSampleStyleSheet()
    buf = io.BytesIO()
    flow = [
        Paragraph("Frame Warranty Terms", styles["Heading2"]),
        Paragraph(
            "Aluminium frames carry a 7 year warranty from the delivery date. " * 6,
            styles["BodyText"],
        ),
        PageBreak(),
        Paragraph("Battery Storage Rules", styles["Heading2"]),
        Paragraph(
            "Store the zorbatron battery at 60 percent charge between 10 and 20 degrees. " * 6,
            styles["BodyText"],
        ),
    ]
    SimpleDocTemplate(buf, pagesize=LETTER).build(flow)
    return buf.getvalue()


def make_docx() -> bytes:
    doc = Document()
    doc.add_heading("Ledger Test Guide", level=0)
    doc.add_heading("Rate Limits", level=1)
    doc.add_paragraph("The quuxplan allows 600 requests per minute with a burst of 50.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Plan", "Limit"
    table.cell(1, 0).text, table.cell(1, 1).text = "Quux", "600 rpm"
    doc.add_heading("Data Export", level=1)
    doc.add_paragraph("Exports are delivered as CSV and expire after 72 hours.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
