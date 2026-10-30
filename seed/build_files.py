"""Build the seed knowledge-base files from seed/corpus/<slug>/*.md.

Each Markdown source becomes a real PDF (reportlab) or DOCX (python-docx) in
seed/files/<slug>/, alternating by file order so both parsers are exercised.
A manifest (seed/files/manifest.json) lists title, file and type for the seeders.

Run:  ai/.venv/bin/python seed/build_files.py
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

SEED = Path(__file__).resolve().parent
CORPUS = SEED / "corpus"
OUT = SEED / "files"
COMPANY = {"northwind": "Northwind Cycles", "ledgerly": "Ledgerly"}


@dataclass
class Block:
    kind: str  # title | h2 | h3 | p | ul | ol | table
    text: str = ""
    items: list | None = None


def parse_markdown(md: str) -> list[Block]:
    blocks: list[Block] = []
    para: list[str] = []
    lst: Block | None = None
    table: list[list[str]] | None = None

    def flush() -> None:
        nonlocal para, lst, table
        if para:
            blocks.append(Block("p", " ".join(para)))
            para = []
        if lst:
            blocks.append(lst)
            lst = None
        if table:
            blocks.append(Block("table", items=table))
            table = None

    for raw in md.splitlines():
        line = raw.rstrip()
        s = line.strip()
        if not s:
            flush()
            continue
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            if table is None:
                flush()
                table = []
            table.append(cells)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", s)
        if m:
            flush()
            blocks.append(
                Block({1: "title", 2: "h2", 3: "h3"}[len(m.group(1))], m.group(2))
            )
            continue
        m = re.match(r"^([-*]|\d+[.)])\s+(.*)$", s)
        if m:
            kind = "ol" if m.group(1)[0].isdigit() else "ul"
            if para:
                blocks.append(Block("p", " ".join(para)))
                para = []
            if lst is None or lst.kind != kind:
                if lst:
                    blocks.append(lst)
                lst = Block(kind, items=[])
            lst.items.append(m.group(2))
            continue
        if lst and raw.startswith(("  ", "\t")):
            lst.items[-1] += " " + s
            continue
        if lst:
            blocks.append(lst)
            lst = None
        para.append(s)
    flush()
    return blocks


def plain(text: str) -> str:
    return re.sub(r"(\*\*|__|`)", "", text)


def rl_markup(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    return re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", text)


def build_pdf(blocks: list[Block], company: str, path: Path) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10.5, leading=14)
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], spaceBefore=10)
    flow = []
    title = next((b.text for b in blocks if b.kind == "title"), path.stem)
    for b in blocks:
        if b.kind == "title":
            flow += [Paragraph(rl_markup(b.text), styles["Title"]), Spacer(1, 6)]
        elif b.kind == "h2":
            flow.append(Paragraph(rl_markup(b.text), h2))
        elif b.kind == "h3":
            flow.append(Paragraph(rl_markup(b.text), styles["Heading3"]))
        elif b.kind == "p":
            flow.append(Paragraph(rl_markup(b.text), body))
        elif b.kind in ("ul", "ol"):
            items = [
                ListItem(Paragraph(rl_markup(t), body), leftIndent=12) for t in b.items
            ]
            flow.append(
                ListFlowable(
                    items,
                    bulletType="1" if b.kind == "ol" else "bullet",
                    start="1" if b.kind == "ol" else None,
                    leftIndent=14,
                )
            )
        elif b.kind == "table":
            data = [[Paragraph(rl_markup(c), body) for c in row] for row in b.items]
            t = Table(data, hAlign="LEFT", repeatRows=1)
            t.setStyle(
                TableStyle(
                    [
                        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ]
                )
            )
            flow += [Spacer(1, 4), t, Spacer(1, 4)]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(
            inch * 0.75,
            inch * 0.5,
            f"{company} (fictional company, demo data) · {title}",
        )
        canvas.drawRightString(LETTER[0] - inch * 0.75, inch * 0.5, f"Page {doc.page}")
        canvas.restoreState()

    SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        title=title,
        author=company,
        leftMargin=inch * 0.8,
        rightMargin=inch * 0.8,
        topMargin=inch * 0.8,
        bottomMargin=inch * 0.8,
    ).build(flow, onFirstPage=footer, onLaterPages=footer)


def build_docx(blocks: list[Block], company: str, path: Path) -> None:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.core_properties.author = company
    for b in blocks:
        if b.kind == "title":
            doc.add_heading(plain(b.text), level=0)
            doc.core_properties.title = plain(b.text)
        elif b.kind == "h2":
            doc.add_heading(plain(b.text), level=1)
        elif b.kind == "h3":
            doc.add_heading(plain(b.text), level=2)
        elif b.kind == "p":
            doc.add_paragraph(plain(b.text))
        elif b.kind in ("ul", "ol"):
            style = "List Number" if b.kind == "ol" else "List Bullet"
            for item in b.items:
                doc.add_paragraph(plain(item), style=style)
        elif b.kind == "table":
            rows = b.items
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Light Grid Accent 1"
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = plain(cell)
    doc.save(str(path))


def main() -> None:
    manifest = []
    for slug_dir in sorted(p for p in CORPUS.iterdir() if p.is_dir()):
        slug = slug_dir.name
        company = COMPANY.get(slug, slug.title())
        out_dir = OUT / slug
        out_dir.mkdir(parents=True, exist_ok=True)
        for i, src in enumerate(sorted(slug_dir.glob("*.md"))):
            blocks = parse_markdown(src.read_text(encoding="utf-8"))
            title = plain(next((b.text for b in blocks if b.kind == "title"), src.stem))
            kind = "pdf" if i % 2 == 0 else "docx"
            target = out_dir / f"{src.stem}.{kind}"
            (build_pdf if kind == "pdf" else build_docx)(blocks, company, target)
            manifest.append(
                {
                    "tenant": slug,
                    "doc": src.stem,
                    "title": title,
                    "source_type": kind,
                    "file": str(target.relative_to(SEED)),
                    "mime": "application/pdf"
                    if kind == "pdf"
                    else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                }
            )
            print(f"{target.relative_to(SEED.parent)}  ({title})")
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(f"{len(manifest)} files, manifest at seed/files/manifest.json")


if __name__ == "__main__":
    main()
